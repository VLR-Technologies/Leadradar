from io import BytesIO
from typing import Any
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.business import BusinessResponse

HEADERS = (
    "S.No",
    "Business Name",
    "Category",
    "Subcategory",
    "Phone",
    "Alternate Phones",
    "Email",
    "Alternate Emails",
    "Official Website",
    "Website Status",
    "Website Type",
    "Directory Links",
    "WhatsApp",
    "Alternate WhatsApp",
    "Address",
    "Locality",
    "City",
    "District",
    "State",
    "PIN Code",
    "Country",
    "Latitude",
    "Longitude",
    "Social Links",
    "Rating",
    "Review Count",
    "Rating Source",
    "Lead Score",
    "Opportunity",
    "Opportunity Reasons",
    "Discovery Sources",
    "Phone Source",
    "Email Source",
    "Website Source",
    "Confidence",
    "Enrichment Status",
    "Scraped At",
)


def _safe_text(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    cleaned = value.replace("\x00", "").replace("\r\n", "\n").strip()
    if cleaned.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{cleaned}"
    return cleaned


def _alternatives(primary: str | None, values: list[str]) -> str:
    primary_key = (primary or "").casefold()
    return "; ".join(value for value in values if value.casefold() != primary_key)


def _source(business: BusinessResponse, field_name: str) -> str:
    values = business.field_provenance.get(field_name, [])
    return values[0].source if values else ""


def _row(index: int, business: BusinessResponse) -> tuple[Any, ...]:
    return (
        index,
        business.name,
        business.category or "",
        "; ".join(business.subcategories),
        business.phone or "",
        _alternatives(business.phone, business.phones),
        business.email or "",
        _alternatives(business.email, business.emails),
        business.website or "",
        business.website_status,
        business.website_type,
        "\n".join(business.directory_links),
        business.whatsapp_number or "",
        _alternatives(business.whatsapp_number, business.whatsapp_numbers),
        business.address.formatted or "",
        business.address.locality or "",
        business.address.city or "",
        business.address.district or "",
        business.address.state or "",
        business.address.postcode or "",
        business.address.country or "",
        business.latitude,
        business.longitude,
        "\n".join(business.social_links),
        business.rating,
        business.review_count,
        business.rating_source or "",
        business.lead_score,
        business.opportunity_level,
        "; ".join(business.opportunity_reasons),
        ", ".join(business.sources),
        _source(business, "phone"),
        _source(business, "email"),
        _source(business, "website"),
        business.confidence,
        business.enrichment_status,
        business.scraped_at.isoformat() if business.scraped_at else "",
    )


def _set_hyperlink(cell: Cell, value: str | None) -> None:
    if not value:
        return
    try:
        parsed = urlsplit(value if "://" in value else f"https://{value}")
    except ValueError:
        return
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return
    cell.hyperlink = parsed.geturl()
    cell.style = "Hyperlink"


def build_leads_workbook(businesses: list[BusinessResponse]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "LeadRadar Leads"
    worksheet.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="177454")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 24

    website_column = HEADERS.index("Official Website") + 1
    for index, business in enumerate(businesses, start=1):
        worksheet.append(tuple(_safe_text(value) for value in _row(index, business)))
        _set_hyperlink(worksheet.cell(row=index + 1, column=website_column), business.website)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_index, header in enumerate(HEADERS, start=1):
        values = [str(header)]
        values.extend(
            str(worksheet.cell(row=row, column=column_index).value or "")
            for row in range(2, min(worksheet.max_row, 100) + 1)
        )
        width = min(45, max(10, max(len(line) for value in values for line in value.splitlines()) + 2))
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
