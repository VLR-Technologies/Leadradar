from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.schemas.business import BusinessResponse
from app.services.excel_export import HEADERS, build_leads_workbook


def lead(*, name: str = "ABC Dental Clinic") -> BusinessResponse:
    return BusinessResponse.model_validate(
        {
            "sourceId": "overture:1",
            "source": "overture",
            "leadId": "lead-1",
            "name": name,
            "category": "Dentist",
            "subcategories": ["cosmetic_dentist"],
            "address": {
                "formatted": "Banjara Hills, Hyderabad, Telangana 500034",
                "locality": "Banjara Hills",
                "city": "Hyderabad",
                "state": "Telangana",
                "postcode": "500034",
                "country": "India",
            },
            "latitude": 17.4,
            "longitude": 78.4,
            "phone": "+91 98765 43210",
            "phones": ["+91 98765 43210", "040-12345678"],
            "normalizedPhones": ["+919876543210", "+914012345678"],
            "email": "hello@abcdental.in",
            "emails": ["hello@abcdental.in"],
            "website": "https://abcdental.in",
            "websites": ["https://abcdental.in"],
            "websiteStatus": "verified",
            "openingHours": None,
            "socialLinks": ["https://instagram.com/abcdental"],
            "sources": ["overture", "openstreetmap"],
            "sourceIds": {"overture": ["overture:1"], "openstreetmap": ["node:2"]},
            "fieldProvenance": {
                "phone": [
                    {
                        "value": "+91 98765 43210",
                        "source": "overture",
                        "confidence": "medium",
                    }
                ],
                "email": [
                    {
                        "value": "hello@abcdental.in",
                        "source": "official_website",
                        "confidence": "high",
                    }
                ],
                "website": [
                    {
                        "value": "https://abcdental.in",
                        "source": "overture",
                        "confidence": "medium",
                    }
                ],
            },
            "leadScore": 85,
            "opportunityLevel": "High",
            "opportunityReasons": ["Public phone available", "No public email found"],
            "enrichmentStatus": "completed",
        }
    )


def test_excel_export_has_calling_columns_and_readable_sheet() -> None:
    payload = build_leads_workbook([lead()])
    workbook = load_workbook(BytesIO(payload))
    worksheet = workbook["LeadRadar Leads"]

    assert tuple(cell.value for cell in worksheet[1]) == HEADERS
    assert worksheet.freeze_panes == "A2"
    final_column = get_column_letter(len(HEADERS))
    assert worksheet.auto_filter.ref == f"A1:{final_column}2"
    assert worksheet["B2"].value == "ABC Dental Clinic"
    assert worksheet["E2"].value == "'+91 98765 43210"
    assert worksheet["F2"].value == "040-12345678"
    website_column = get_column_letter(HEADERS.index("Official Website") + 1)
    opportunity_column = get_column_letter(HEADERS.index("Opportunity") + 1)
    assert worksheet[f"{website_column}2"].hyperlink.target == "https://abcdental.in"
    assert worksheet[f"{opportunity_column}2"].value == "High"


def test_excel_export_neutralizes_formula_injection() -> None:
    payload = build_leads_workbook([lead(name="=HYPERLINK(\"https://bad.example\")")])
    worksheet = load_workbook(BytesIO(payload), data_only=False).active

    assert worksheet["B2"].data_type != "f"
    assert worksheet["B2"].value.startswith("'=")
