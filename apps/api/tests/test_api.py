from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.api.dependencies import get_discovery_service, get_enrichment_service
from app.core.categories import BusinessCategory
from app.main import app
from app.models.business import Business, BusinessAddress
from app.models.enrichment import ExtractedContact, WebsiteInspection
from app.models.location import ResolvedLocation
from app.providers.enrichment.base import WebsiteTimeoutError
from app.providers.enrichment.search import DisabledBusinessSearchProvider
from app.services.business_discovery import BusinessDiscoveryService
from app.services.business_enrichment import BusinessEnrichmentService
from app.services.location_resolver import LocationResolver


class EmptyProvider:
    async def discover(
        self,
        *,
        location: ResolvedLocation,
        category: BusinessCategory,
        limit: int,
    ) -> list[Business]:
        return []


class ManyProvider:
    def __init__(self, count: int) -> None:
        self.count = count

    async def discover(
        self,
        *,
        location: ResolvedLocation,
        category: BusinessCategory,
        limit: int,
    ) -> list[Business]:
        return [
            Business(
                source_id=f"overture:export-{index}",
                source="overture",
                name=f"Export Test Business {index}",
                category=category.label,
                address=BusinessAddress(
                    street=f"{index} Export Road",
                    city=location.city,
                    state=location.region,
                    country=location.country,
                    formatted=(
                        f"{index} Export Road, {location.city}, "
                        f"{location.region}, {location.country}"
                    ),
                ),
                latitude=17.0 + (index * 0.01),
                longitude=78.0 + (index * 0.01),
                phone=f"+91980000{index:04d}",
                email=None,
                website=None,
                opening_hours=None,
            )
            for index in range(min(self.count, limit))
        ]


class VerifiedWebsiteProvider:
    async def inspect(self, *, business: Business, website_url: str) -> WebsiteInspection:
        return WebsiteInspection(
            verification_status="verified",
            final_url="https://clinic.in",
            phones=(ExtractedContact("+919876543210", "high", website_url),),
            emails=(ExtractedContact("hello@clinic.in", "high", website_url),),
            visited_pages=(website_url,),
        )


class TimeoutWebsiteProvider:
    async def inspect(self, *, business: Business, website_url: str) -> WebsiteInspection:
        raise WebsiteTimeoutError


def enrichment_service(website_provider=None) -> BusinessEnrichmentService:
    return BusinessEnrichmentService(
        website_provider=website_provider or VerifiedWebsiteProvider(),
        search_provider=DisabledBusinessSearchProvider(),
        batch_limit=20,
        max_concurrency=2,
    )


def enrichment_business_payload(*, website: str | None = "https://clinic.in") -> dict:
    return {
        "sourceId": "node:123",
        "source": "openstreetmap",
        "name": "ABC Dental Clinic",
        "category": "Dentist",
        "address": {
            "street": None,
            "houseNumber": None,
            "postcode": None,
            "city": "Hyderabad",
            "state": "Telangana",
            "country": "India",
            "formatted": "Hyderabad, Telangana, India",
        },
        "latitude": 17.4,
        "longitude": 78.4,
        "phone": None,
        "email": None,
        "website": website,
        "openingHours": None,
    }


@pytest.mark.asyncio
async def test_health(api_client) -> None:
    response = await api_client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


@pytest.mark.asyncio
async def test_categories_are_returned_without_raw_osm_tags(api_client) -> None:
    response = await api_client.get("/api/v1/categories")

    assert response.status_code == 200
    categories = response.json()["categories"]
    assert {category["label"] for category in categories} >= {
        "Dentist",
        "Doctor / Clinic",
        "Real Estate Agency",
    }
    assert all(set(category) == {"id", "label"} for category in categories)


@pytest.mark.asyncio
async def test_invalid_category_returns_validation_error_without_provider_call(api_client) -> None:
    response = await api_client.post(
        "/api/v1/businesses/discover",
        json={
            "countryCode": "DE",
            "region": None,
            "city": "Berlin",
            "category": "Made Up Category",
            "limit": 20,
        },
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "Unsupported business category: Made Up Category"


@pytest.mark.asyncio
@pytest.mark.parametrize("limit", [0, 501])
async def test_result_limit_validation(api_client, limit: int) -> None:
    response = await api_client.post(
        "/api/v1/businesses/discover",
        json={
            "countryCode": "DE",
            "region": None,
            "city": "Berlin",
            "category": "Dentist",
            "limit": limit,
        },
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_location_country_endpoint(api_client) -> None:
    response = await api_client.get("/api/v1/locations/countries")

    assert response.status_code == 200
    assert response.json() == [
        {"code": "DE", "name": "Germany", "requiresRegion": False},
        {"code": "GB", "name": "United Kingdom", "requiresRegion": False},
        {"code": "US", "name": "United States", "requiresRegion": True},
        {"code": "IN", "name": "India", "requiresRegion": False},
    ]


@pytest.mark.asyncio
async def test_location_city_and_region_endpoints(api_client) -> None:
    region_response = await api_client.get(
        "/api/v1/locations/regions",
        params={"countryCode": "US"},
    )
    city_response = await api_client.get(
        "/api/v1/locations/cities",
        params={"countryCode": "US", "region": "California"},
    )

    assert region_response.status_code == 200
    assert {item["name"] for item in region_response.json()} >= {
        "California",
        "Texas",
        "Washington",
    }
    assert city_response.status_code == 200
    assert {item["name"] for item in city_response.json()} >= {
        "Los Angeles",
        "San Diego",
        "San Francisco",
    }
    assert all(item["region"] == "California" for item in city_response.json())


@pytest.mark.asyncio
async def test_unsupported_location_endpoint_country_is_rejected(api_client) -> None:
    response = await api_client.get(
        "/api/v1/locations/cities",
        params={"countryCode": "ZZ"},
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_us_discovery_requires_region(api_client) -> None:
    response = await api_client.post(
        "/api/v1/businesses/discover",
        json={
            "countryCode": "US",
            "region": None,
            "city": "San Francisco",
            "category": "Cafe",
            "limit": 5,
        },
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "Region is required for United States searches."


@pytest.mark.asyncio
async def test_discovery_uses_new_api_schema_and_resolved_metadata(api_client) -> None:
    service = BusinessDiscoveryService(
        provider=EmptyProvider(),
        location_resolver=LocationResolver(),
    )
    app.dependency_overrides[get_discovery_service] = lambda: service
    try:
        response = await api_client.post(
            "/api/v1/businesses/discover",
            json={
                "countryCode": "US",
                "region": "California",
                "city": "San Francisco",
                "category": "Cafe",
                "limit": 5,
            },
        )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.json()["query"] == {
        "countryCode": "US",
        "country": "United States",
        "region": "California",
        "city": "San Francisco",
        "category": "Cafe",
        "limit": 5,
    }
    assert response.json()["count"] == 0
    assert response.json()["totalCount"] == 0
    assert response.json()["pageSize"] == 50
    assert response.json()["summary"]["total"] == 0
    session_id = response.json()["sessionId"]
    page_response = await api_client.get(
        f"/api/v1/businesses/search-sessions/{session_id}",
        params={"page": 1, "pageSize": 25, "filter": "phone"},
    )
    assert page_response.status_code == 200
    assert page_response.json()["sessionId"] == session_id
    assert page_response.json()["businesses"] == []


@pytest.mark.asyncio
async def test_all_available_limit_is_bounded_and_reported(api_client) -> None:
    service = BusinessDiscoveryService(
        provider=EmptyProvider(),
        location_resolver=LocationResolver(),
    )
    app.dependency_overrides[get_discovery_service] = lambda: service
    try:
        response = await api_client.post(
            "/api/v1/businesses/discover",
            json={
                "countryCode": "IN",
                "region": "Telangana",
                "city": "Hyderabad",
                "category": "Restaurant",
                "limit": "all",
            },
        )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.json()["query"]["limit"] == "all"
    assert response.json()["effectiveLimit"] == 2_000


@pytest.mark.asyncio
async def test_single_enrichment_request_returns_provenance(api_client) -> None:
    app.dependency_overrides[get_enrichment_service] = enrichment_service
    try:
        response = await api_client.post(
            "/api/v1/businesses/enrich",
            json={"business": enrichment_business_payload()},
        )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "completed"
    assert payload["websiteStatus"] == "verified"
    assert payload["reasonCode"] == "ADDITIONAL_CONTACTS_FOUND"
    assert payload["fields"]["phone"]["primary"] == {
        "value": "+919876543210",
        "source": "official_website",
        "confidence": "high",
        "sourceUrl": "https://clinic.in",
        "sourceType": "visible_text",
    }


@pytest.mark.asyncio
async def test_invalid_enrichment_request_is_rejected(api_client) -> None:
    response = await api_client.post(
        "/api/v1/businesses/enrich",
        json={"business": {"sourceId": "node:123"}},
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_business_without_website_returns_structured_result(api_client) -> None:
    app.dependency_overrides[get_enrichment_service] = enrichment_service
    try:
        response = await api_client.post(
            "/api/v1/businesses/enrich",
            json={"business": enrichment_business_payload(website=None)},
        )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.json()["status"] == "no_data"
    assert response.json()["reasonCode"] == "SEARCH_PROVIDER_DISABLED"
    assert response.json()["message"] == (
        "No official website is available from the current sources."
    )


@pytest.mark.asyncio
async def test_website_timeout_is_returned_without_losing_business(api_client) -> None:
    app.dependency_overrides[get_enrichment_service] = lambda: enrichment_service(
        TimeoutWebsiteProvider()
    )
    try:
        response = await api_client.post(
            "/api/v1/businesses/enrich",
            json={"business": enrichment_business_payload()},
        )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.json()["status"] == "failed"
    assert response.json()["websiteStatus"] == "unreachable"
    assert response.json()["reasonCode"] == "WEBSITE_TIMEOUT"
    assert response.json()["websiteAudit"] == {
        "reachable": False,
        "usesHttps": None,
        "redirectBehavior": None,
        "mobileViewport": None,
        "contactPagePresent": None,
        "emailPresent": None,
        "phonePresent": None,
        "socialLinksPresent": None,
        "titlePresent": None,
        "metaDescriptionPresent": None,
        "brokenResponseCount": 0,
        "label": "Website unreachable",
    }


@pytest.mark.asyncio
async def test_batch_enrichment_rejects_more_than_twenty_businesses(api_client) -> None:
    response = await api_client.post(
        "/api/v1/businesses/enrich-batch",
        json={"businesses": [enrichment_business_payload() for _ in range(21)]},
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_excel_export_endpoint_returns_downloadable_workbook(api_client) -> None:
    response = await api_client.post(
        "/api/v1/businesses/export",
        json={
            "businesses": [enrichment_business_payload()],
            "searchLabel": "Hyderabad Dentist",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "leadradar-hyderabad-dentist-" in response.headers["content-disposition"]
    assert response.content.startswith(b"PK")


@pytest.mark.asyncio
async def test_excel_export_rejects_empty_results(api_client) -> None:
    response = await api_client.post(
        "/api/v1/businesses/export",
        json={"businesses": []},
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_session_export_contains_all_rows_not_only_current_page(api_client) -> None:
    service = BusinessDiscoveryService(
        provider=ManyProvider(60),
        location_resolver=LocationResolver(),
    )
    app.dependency_overrides[get_discovery_service] = lambda: service
    try:
        discovery_response = await api_client.post(
            "/api/v1/businesses/discover",
            json={
                "countryCode": "IN",
                "region": "Telangana",
                "city": "Hyderabad",
                "category": "Restaurant",
                "limit": 100,
                "pageSize": 25,
            },
        )
        session_id = discovery_response.json()["sessionId"]
        export_response = await api_client.get(
            f"/api/v1/businesses/search-sessions/{session_id}/export",
            params={"filter": "all"},
        )
    finally:
        app.dependency_overrides.clear()

    assert discovery_response.status_code == 200
    assert discovery_response.json()["count"] == 25
    assert discovery_response.json()["totalCount"] == 60
    assert export_response.status_code == 200
    worksheet = load_workbook(BytesIO(export_response.content), read_only=True).active
    assert worksheet.max_row == 61
